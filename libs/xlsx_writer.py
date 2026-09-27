"""Запись перечня НТД в XLSX на основе шаблона «Перечень_НТД.xlsx».

Шаблон содержит один лист «Лист1» с шапкой из 5 колонок:
    A: №
    B: НТД
    C: Кол-во упоминаний
    D: Номера листов файла pdf
    E: Название файла
Строка 2 в шаблоне — эталон формата (границы, шрифт, выравнивание),
её стили копируются во все добавляемые строки. Автофильтр закреплён
на весь диапазон таблицы A1:E{последняя_строка}.
"""

from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

# Количество колонок и порядок данных в шаблоне
TEMPLATE_COLS = 5
START_ROW = 2  # первая строка — шапка, данные начинаются со второй


def _copy_row_format(ws: Worksheet, src_row: int, dst_row: int, cols: int) -> None:
    """Копирует формат ячеек строки src_row в строку dst_row.

    Копируется весь стиль ячейки (шрифт, границы, заливка, выравнивание,
    числовой формат) через внутренний атрибут `_style` — стандартный приём
    openpyxl для полного переноса форматирования. Значения ячеек
    не затрагиваются.
    """
    for col in range(1, cols + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        dst._style = copy(src._style)


def analysis_to_xlsx_bytes(
    rows: list[tuple],
    template_path: str | Path,
) -> bytes:
    """Формирует XLSX-файл перечня НТД на основе шаблона.

    Аргументы:
        rows — список кортежей (номер, нтд, кол_во, листы, имя_файла);
        template_path — путь к файлу-шаблону «Перечень_НТД.xlsx».

    Возвращает содержимое XLSX-файла в виде байтов (удобно для отдачи
    через Streamlit `st.download_button` или записи на диск).
    """
    template = Path(template_path)
    if not template.exists():
        raise FileNotFoundError(f"Шаблон не найден: {template}")

    wb = load_workbook(template)
    ws = wb["Лист1"]

    # Данные начинаются со второй строки (первая — шапка таблицы)
    for offset, row in enumerate(rows):
        row_idx = START_ROW + offset
        for col, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col).value = value
        # Для третьей и следующих строк копируем формат из второй строки
        if row_idx > START_ROW:
            _copy_row_format(ws, START_ROW, row_idx, TEMPLATE_COLS)

    # Фиксируем автофильтр на всю таблицу, чтобы он покрывал новые строки
    last_row = max(START_ROW, START_ROW + len(rows) - 1)
    ws.auto_filter.ref = f"A1:E{last_row}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()